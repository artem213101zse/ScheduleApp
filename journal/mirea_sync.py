"""Download and parse official MIREA Excel timetable into journal lessons."""

from __future__ import annotations

import re
import ssl
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from .models import Group, JournalSettings, Lesson, StudentProfile, Subject, Teacher

GROUP_RE = re.compile(r"([А-Яа-яA-Za-z]{4}-\d{2}-\d{2})")
HREF_RE = re.compile(r"""href\s*=\s*["']([^"']+\.(?:xlsx?|XLSX?))["']""", re.I)
WEEK_HINT_RE = re.compile(r"кр\.?\s*([\d,\s\-]+)\s*н", re.I)
UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)
WEEKDAY_NAMES = {
    "понедельник": 0,
    "пн": 0,
    "вторник": 1,
    "вт": 1,
    "среда": 2,
    "ср": 2,
    "четверг": 3,
    "чт": 3,
    "пятница": 4,
    "пт": 4,
    "суббота": 5,
    "сб": 5,
    "воскресенье": 6,
    "вс": 6,
}


class SyncError(Exception):
    pass


@dataclass
class WeeklyLesson:
    weekday: int
    start: time
    end: time
    parity: str
    weeks: list[int] | None
    subject: str
    lesson_type: str
    teacher: str
    room: str


@dataclass
class ConcreteLesson:
    date: date
    start: time
    end: time
    subject: str
    lesson_type: str
    teacher: str
    room: str


def normalize_group(name: str) -> str:
    return re.sub(r"\s+", "", name or "").upper().replace("Ё", "Е")


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def parse_time_value(value) -> time | None:
    if value is None or value == "":
        return None
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        total = int(round(float(value) * 24 * 60))
        return time((total // 60) % 24, total % 60)
    text = cell_text(value).replace(".", ":").replace("—", "-").replace("–", "-")
    match = re.search(r"(\d{1,2})[:\-](\d{2})", text)
    if not match:
        return None
    return time(int(match.group(1)) % 24, int(match.group(2)) % 60)


def parse_weekday(value) -> int | None:
    text = cell_text(value).lower().replace("ё", "е")
    if not text:
        return None
    first = text.split()[0]
    return WEEKDAY_NAMES.get(first)


def parse_parity(value) -> str | None:
    text = cell_text(value).upper().replace(" ", "")
    if text in {"I", "1", "ЧИСЛ", "НЧ"}:
        return "odd"
    if text in {"II", "2", "ЗНАМ", "ЧЕТ"}:
        return "even"
    return None


def parse_week_list(text: str) -> list[int] | None:
    match = WEEK_HINT_RE.search(text or "")
    if not match:
        return None
    weeks: list[int] = []
    for part in match.group(1).split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            left, right = part.split("-", 1)
            if left.strip().isdigit() and right.strip().isdigit():
                weeks.extend(range(int(left), int(right) + 1))
        elif part.isdigit():
            weeks.append(int(part))
    return weeks or None


def clean_subject(text: str) -> str:
    text = WEEK_HINT_RE.sub("", text or "")
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ,;.-")


def map_lesson_type(text: str) -> str:
    lowered = (text or "").lower()
    if "лек" in lowered:
        return "ЛК"
    if "лаб" in lowered:
        return "ЛБ"
    if "сем" in lowered or "с/р" in lowered or "с.р" in lowered:
        return "СМ"
    return "ПР"


def short_teacher(name: str) -> str:
    parts = [p for p in re.split(r"\s+", (name or "").strip()) if p]
    if len(parts) >= 3:
        return f"{parts[0]} {parts[1][0]}. {parts[2][0]}."
    if len(parts) == 2 and len(parts[1]) > 1:
        return f"{parts[0]} {parts[1][0]}."
    return name.strip() or "—"


def _fill_merges(grid: list[list], merges) -> None:
    for item in merges:
        if hasattr(item, "min_row"):
            r0, r1, c0, c1 = item.min_row, item.max_row, item.min_col, item.max_col
        else:
            r0, r1, c0, c1 = item
            r0 += 1
            r1 += 1
            c0 += 1
            c1 += 1
        val = grid[r0 - 1][c0 - 1]
        for row in range(r0, r1 + 1):
            for col in range(c0, c1 + 1):
                if grid[row - 1][col - 1] in (None, ""):
                    grid[row - 1][col - 1] = val


def load_sheets(path: str | Path) -> list[list[list]]:
    path = str(path)
    lower = path.lower()
    if lower.endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(path)
        sheets = []
        for sheet in book.sheets():
            grid = []
            for r in range(sheet.nrows):
                grid.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
            merges = getattr(sheet, "merged_cells", []) or []
            _fill_merges(grid, merges)
            sheets.append(grid)
        return sheets

    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        max_r = ws.max_row or 0
        max_c = ws.max_column or 0
        grid = [[None] * max_c for _ in range(max_r)]
        for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c):
            for cell in row:
                grid[cell.row - 1][cell.column - 1] = cell.value
        _fill_merges(grid, ws.merged_cells.ranges)
        sheets.append(grid)
    return sheets


def _find_group_cell(grid: list[list], group_name: str) -> tuple[int, int] | None:
    want = normalize_group(group_name)
    for r, row in enumerate(grid[:40]):
        for c, value in enumerate(row):
            text = normalize_group(cell_text(value))
            found = GROUP_RE.search(text)
            if found and normalize_group(found.group(1)) == want:
                return r, c
            if text == want:
                return r, c
    return None


def parse_excel(path: str | Path, group_name: str) -> list[WeeklyLesson]:
    want = normalize_group(group_name)
    if not want:
        raise SyncError("Не указана учебная группа для синхронизации")
    lessons: list[WeeklyLesson] = []
    for grid in load_sheets(path):
        if not grid:
            continue
        found = _find_group_cell(grid, group_name)
        if found is None:
            continue
        row_i, col_i = found
        start_row = row_i + 2
        weekday = num = start = end = parity = None
        for r in range(start_row, min(len(grid), start_row + 120)):
            row = grid[r]
            def at(offset):
                idx = col_i + offset
                if idx < 0 or idx >= len(row):
                    return None
                return row[idx]

            weekday = parse_weekday(at(-5)) if parse_weekday(at(-5)) is not None else weekday
            if cell_text(at(-4)):
                try:
                    num = int(float(cell_text(at(-4)).split()[0]))
                except (TypeError, ValueError):
                    pass
            start = parse_time_value(at(-3)) or start
            end = parse_time_value(at(-2)) or end
            parity = parse_parity(at(-1)) or parity
            subject_raw = cell_text(at(0))
            if weekday is None or not start or not end or not subject_raw:
                continue
            if GROUP_RE.search(subject_raw) or normalize_group(subject_raw) == want:
                continue
            subject = clean_subject(subject_raw)
            if subject.lower() in {"предмет", "дисциплина", "видзанятий"}:
                continue
            if subject:
                lessons.append(
                    WeeklyLesson(
                        weekday=weekday,
                        start=start,
                        end=end,
                        parity=parity or "odd",
                        weeks=parse_week_list(subject_raw),
                        subject=subject,
                        lesson_type=map_lesson_type(cell_text(at(1)) or subject_raw),
                        teacher=short_teacher(cell_text(at(2))),
                        room=cell_text(at(3))[:32],
                    )
                )
    if not lessons:
        raise SyncError(f"В файле расписания нет пар для группы {group_name}")
    return lessons


def monday_of(day: date) -> date:
    return day - timedelta(days=day.weekday())


def academic_week(day: date, semester_start: date) -> int:
    return ((monday_of(day) - monday_of(semester_start)).days // 7) + 1


def semester_bounds(cfg: JournalSettings | None = None) -> tuple[date, date]:
    cfg = cfg or JournalSettings.get()
    if cfg.semester_start and cfg.semester_end:
        return cfg.semester_start, cfg.semester_end
    today = timezone.localdate()
    if today.month >= 8:
        return date(today.year, 9, 1), date(today.year, 12, 31)
    return date(today.year, 2, 9), date(today.year, 6, 30)


def expand_semester(weekly: list[WeeklyLesson], start: date, end: date) -> list[ConcreteLesson]:
    out: list[ConcreteLesson] = []
    day = start
    while day <= end:
        week_no = academic_week(day, start)
        odd = week_no % 2 == 1
        for item in weekly:
            if item.weekday != day.weekday():
                continue
            if item.weeks is not None:
                if week_no not in item.weeks:
                    continue
            elif item.parity == "odd" and not odd:
                continue
            elif item.parity == "even" and odd:
                continue
            out.append(
                ConcreteLesson(
                    date=day,
                    start=item.start,
                    end=item.end,
                    subject=item.subject,
                    lesson_type=item.lesson_type,
                    teacher=item.teacher,
                    room=item.room,
                )
            )
        day += timedelta(days=1)
    return out


def resolve_excel_group_name() -> str:
    override = (getattr(settings, "MIREA_SCHEDULE_GROUP", "") or "").strip()
    if override:
        return override
    cfg = JournalSettings.get()
    if cfg.sync_group_name.strip():
        return cfg.sync_group_name.strip()
    student = StudentProfile.objects.select_related("group").first()
    if student and student.group_id:
        return student.group.name
    group = Group.objects.filter(is_journal=True).first()
    if group:
        return group.name
    raise SyncError("Не найдена группа для синхронизации")


def _http_get(url: str, timeout: int = 45) -> bytes:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": UA,
            "Accept": "*/*",
            "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.8",
        },
    )
    context = ssl.create_default_context()
    try:
        with urllib.request.urlopen(req, timeout=timeout, context=context) as resp:
            return resp.read()
    except ssl.SSLError:
        context = ssl._create_unverified_context()
        with urllib.request.urlopen(req, timeout=timeout, context=context) as resp:
            return resp.read()


def _schedule_links(html: str, base: str) -> list[str]:
    links = []
    seen = set()
    for href in HREF_RE.findall(html or ""):
        url = urllib.parse.urljoin(base, href.replace("&amp;", "&"))
        if url not in seen:
            seen.add(url)
            links.append(url)
    return links


def _link_score(url: str, group_name: str) -> int:
    lowered = urllib.parse.unquote(url).lower()
    score = 0
    compact = normalize_group(group_name).lower()
    if compact and compact.replace("-", "") in lowered.replace("-", ""):
        score += 100
    prefix = compact[:4]
    if prefix and prefix.lower() in lowered:
        score += 40
    if any(word in lowered for word in ("магистр", "magistr", "маг_")):
        score += 25
    if any(word in lowered for word in ("вечер", "заоч", "очно-заоч")):
        score += 15
    if any(word in lowered for word in ("экзам", "зачет", "зачёт", "session", "сессия")):
        score -= 50
    return score


def download_excel_for_group(group_name: str) -> Path:
    local = (getattr(settings, "MIREA_SCHEDULE_FILE", "") or "").strip()
    if local:
        path = Path(local)
        if not path.is_file():
            raise SyncError(f"Локальный файл расписания не найден: {local}")
        return path

    cfg = JournalSettings.get()
    page_url = getattr(settings, "MIREA_SCHEDULE_URL", "https://www.mirea.ru/schedule/")
    candidates: list[str] = []
    if cfg.sync_file_url:
        candidates.append(cfg.sync_file_url)
    try:
        html = _http_get(page_url).decode("utf-8", "ignore")
    except (urllib.error.URLError, TimeoutError, ssl.SSLError) as exc:
        raise SyncError(
            "Не удалось открыть страницу расписания МИРЭА. "
            "На сервере должна быть сеть до www.mirea.ru."
        ) from exc
    links = sorted(_schedule_links(html, page_url), key=lambda u: -_link_score(u, group_name))
    for url in links:
        if url not in candidates:
            candidates.append(url)
    if not candidates:
        raise SyncError("На странице расписания МИРЭА нет ссылок на Excel")

    last_error = None
    tmpdir = Path(tempfile.mkdtemp(prefix="mirea-sched-"))
    for url in candidates[:18]:
        try:
            data = _http_get(url)
        except (urllib.error.URLError, TimeoutError, ssl.SSLError) as exc:
            last_error = exc
            continue
        if len(data) < 1000:
            continue
        suffix = ".xlsx" if data[:2] == b"PK" else ".xls"
        path = tmpdir / f"schedule{suffix}"
        path.write_bytes(data)
        try:
            parse_excel(path, group_name)
        except SyncError:
            continue
        cfg.sync_file_url = url
        cfg.save(update_fields=["sync_file_url"])
        return path
    raise SyncError(
        f"Группа {group_name} не найдена в Excel на mirea.ru/schedule. "
        + (f"({last_error})" if last_error else "")
    )


def _get_subject(name: str) -> Subject:
    name = clean_subject(name) or "Дисциплина"
    obj, _ = Subject.objects.get_or_create(name=name[:160])
    return obj


def _get_teacher(name: str) -> Teacher:
    short = short_teacher(name)[:64]
    full = (name or short).strip()[:120] or short
    obj = Teacher.objects.filter(short_name=short).first()
    if obj:
        return obj
    obj = Teacher.objects.filter(full_name=full).first()
    if obj:
        return obj
    return Teacher.objects.create(short_name=short, full_name=full)


def apply_sync(group: Group, concrete: list[ConcreteLesson]) -> dict:
    created = 0
    updated = 0
    skipped_manual = 0
    kept_pending: set[int] = set()
    with transaction.atomic():
        manuals = {
            (item.date, item.start_time)
            for item in group.lessons.filter(source="manual").only("date", "start_time")
        }
        for item in concrete:
            key = (item.date, item.start)
            if key in manuals:
                skipped_manual += 1
                continue
            subject = _get_subject(item.subject)
            teacher = _get_teacher(item.teacher)
            subject.teachers.add(teacher)
            existing = (
                group.lessons.filter(source="sync", date=item.date, start_time=item.start)
                .order_by("id")
                .first()
            )
            if existing:
                existing.subject = subject
                existing.teacher = teacher
                existing.end_time = item.end
                existing.lesson_type = item.lesson_type
                existing.room = item.room[:32]
                existing.save(
                    update_fields=["subject", "teacher", "end_time", "lesson_type", "room"]
                )
                if not existing.accepted:
                    kept_pending.add(existing.id)
                updated += 1
                continue
            lesson = Lesson.objects.create(
                group=group,
                subject=subject,
                teacher=teacher,
                date=item.date,
                start_time=item.start,
                end_time=item.end,
                lesson_type=item.lesson_type,
                week_type="Н",
                room=item.room[:32],
                source="sync",
                accepted=False,
            )
            created += 1
            kept_pending.add(lesson.id)
        stale = group.lessons.filter(source="sync", accepted=False).exclude(id__in=kept_pending)
        removed = stale.count()
        stale.delete()
    pending = group.lessons.filter(source="sync", accepted=False).count()
    return {
        "created": created,
        "updated": updated,
        "skipped_manual": skipped_manual,
        "removed": removed,
        "pending": pending,
    }


def sync_journal() -> dict:
    group = Group.objects.filter(is_journal=True).first()
    if group is None:
        raise SyncError("Группа журнала не найдена")
    excel_group = resolve_excel_group_name()
    path = download_excel_for_group(excel_group)
    weekly = parse_excel(path, excel_group)
    cfg = JournalSettings.get()
    start, end = semester_bounds(cfg)
    concrete = expand_semester(weekly, start, end)
    if not concrete:
        raise SyncError("В пределах семестра нет пар из официального файла")
    stats = apply_sync(group, concrete)
    cfg.last_sync_at = timezone.now()
    cfg.save(update_fields=["last_sync_at"])
    stats["group"] = excel_group
    stats["from"] = start.isoformat()
    stats["to"] = end.isoformat()
    stats["weekly"] = len(weekly)
    return stats


def cancel_pending() -> int:
    group = Group.objects.filter(is_journal=True).first()
    if group is None:
        return 0
    deleted, _ = group.lessons.filter(source="sync", accepted=False).delete()
    return deleted


def accept_pending(pk: int | None = None) -> int:
    group = Group.objects.filter(is_journal=True).first()
    if group is None:
        return 0
    qs = group.lessons.filter(source="sync", accepted=False)
    if pk is not None:
        qs = qs.filter(pk=pk)
    return qs.update(accepted=True)
