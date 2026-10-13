from datetime import date, time, timedelta

from django.contrib.auth.models import User
from django.test import Client, TestCase

from journal.models import Group, Lesson, Subject, Teacher


class LessonEditorTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="kononov.a.s@edu.mirea.ru",
            email="kononov.a.s@edu.mirea.ru",
            password="mirea2026",
        )
        self.group = Group.objects.create(slug="is", name="ИС-21-1", is_journal=True)
        self.t1 = Teacher.objects.create(short_name="Лебедев А. С.", full_name="Лебедев Александр Сергеевич")
        self.t2 = Teacher.objects.create(short_name="Морозов Д. П.", full_name="Морозов Дмитрий Павлович")
        self.subj = Subject.objects.create(name="Облачные технологии")
        self.subj.teachers.add(self.t1, self.t2)
        self.client = Client()
        self.client.login(username="kononov.a.s@edu.mirea.ru", password="mirea2026")
        session = self.client.session
        session["schedule_edit"] = True
        session.save()

    def test_index_uses_native_selects(self):
        html = self.client.get("/").content.decode("utf-8")
        self.assertIn('<select class="login-input" id="fSubject" required>', html)
        self.assertIn('<select class="login-input" id="fTeacher" required>', html)
        self.assertIn('id="btnCopyToggle"', html)
        self.assertIn('id="fCopyDate"', html)
        self.assertNotIn('id="subjectList"', html)
        self.assertNotIn('list="subjectList"', html)

    def test_bootstrap_nests_subject_teachers(self):
        data = self.client.get("/api/bootstrap/").json()
        cloud = next(s for s in data["subjects"] if s["name"] == "Облачные технологии")
        names = {t["short_name"] for t in cloud["teachers"]}
        self.assertEqual(names, {"Лебедев А. С.", "Морозов Д. П."})

    def test_create_binds_teacher_to_subject(self):
        extra = Teacher.objects.create(short_name="Ким А. В.", full_name="Ким Анна Викторовна")
        resp = self.client.post(
            "/api/lessons/create/",
            data={
                "date": "2026-09-14",
                "subject_id": self.subj.id,
                "teacher_id": extra.id,
                "start": "18:00",
                "end": "19:30",
                "type": "ПР",
                "week": "Н",
                "room": "402",
            },
            content_type="application/json",
        )
        self.assertTrue(resp.json().get("ok"), resp.json())
        self.assertTrue(self.subj.teachers.filter(pk=extra.pk).exists())
        self.assertEqual(Lesson.objects.filter(date=date(2026, 9, 14)).count(), 1)

    def test_copy_creates_same_pair_on_other_day(self):
        payload = {
            "date": "2026-09-07",
            "subject_id": self.subj.id,
            "teacher_id": self.t1.id,
            "start": "18:00",
            "end": "19:30",
            "type": "ПР",
            "week": "Н",
            "room": "402",
        }
        first = self.client.post("/api/lessons/create/", data=payload, content_type="application/json")
        self.assertTrue(first.json().get("ok"), first.json())
        payload["date"] = "2026-09-14"
        payload["start"] = "19:40"
        payload["end"] = "21:10"
        copied = self.client.post("/api/lessons/create/", data=payload, content_type="application/json")
        self.assertTrue(copied.json().get("ok"), copied.json())
        days = list(Lesson.objects.order_by("date", "start_time").values_list("date", "start_time"))
        self.assertEqual(len(days), 2)
        self.assertEqual(days[0][0], date(2026, 9, 7))
        self.assertEqual(days[1][0], date(2026, 9, 14))
        self.assertEqual(days[1][1], time(19, 40))

    def _lesson(self, day, start=None, end=None, week="Н"):
        return Lesson.objects.create(
            group=self.group,
            subject=self.subj,
            teacher=self.t1,
            date=day,
            start_time=start or time(18, 0),
            end_time=end or time(19, 30),
            lesson_type="ПР",
            week_type=week,
            room="402",
        )

    def test_index_has_week_copy_controls(self):
        html = self.client.get("/").content.decode("utf-8")
        self.assertIn('id="btnCopyWeekToggle"', html)
        self.assertIn('id="fCopyWeekDate"', html)
        self.assertIn('id="btnCopyWeekGo"', html)
        self.assertIn('id="autoAttendSwitch"', html)
        self.assertIn('id="btnMarkPast"', html)
        self.assertIn('id="btnSyncSchedule"', html)
        self.assertIn('id="btnAcceptSync"', html)
        self.assertIn('id="btnCancelSync"', html)

    def test_quick_attendance_toggles_mark(self):
        lesson = self._lesson(date(2026, 9, 7), week="Н")
        resp = self.client.post(
            f"/api/lessons/{lesson.pk}/update/",
            data={"week": "+"},
            content_type="application/json",
        )
        self.assertTrue(resp.json().get("ok"), resp.json())
        lesson.refresh_from_db()
        self.assertEqual(lesson.week_type, "+")
        resp = self.client.post(
            f"/api/lessons/{lesson.pk}/update/",
            data={"week": "Н"},
            content_type="application/json",
        )
        self.assertTrue(resp.json().get("ok"), resp.json())
        lesson.refresh_from_db()
        self.assertEqual(lesson.week_type, "Н")

    def test_copy_week_shifts_lessons_and_resets_mark(self):
        self._lesson(date(2026, 9, 7), week="+")
        self._lesson(date(2026, 9, 9), start=time(19, 40), end=time(21, 10), week="+")
        resp = self.client.post(
            "/api/lessons/copy-week/",
            data={"from": "2026-09-08", "to": "2026-09-16"},
            content_type="application/json",
        )
        body = resp.json()
        self.assertTrue(body.get("ok"), body)
        self.assertEqual(body["created"], 2)
        self.assertEqual(body["skipped"], 0)
        monday = Lesson.objects.get(date=date(2026, 9, 14))
        wednesday = Lesson.objects.get(date=date(2026, 9, 16))
        self.assertEqual(monday.week_type, "Н")
        self.assertEqual(monday.subject_id, self.subj.id)
        self.assertEqual(monday.teacher_id, self.t1.id)
        self.assertEqual(monday.room, "402")
        self.assertEqual(wednesday.start_time, time(19, 40))
        again = self.client.post(
            "/api/lessons/copy-week/",
            data={"from": "2026-09-07", "to": "2026-09-14"},
            content_type="application/json",
        )
        self.assertEqual(again.json()["created"], 0)
        self.assertEqual(again.json()["skipped"], 2)

    def test_mark_past_and_auto_attend(self):
        from django.utils import timezone

        today = timezone.localdate()
        yesterday = today - timedelta(days=1)
        tomorrow = today + timedelta(days=1)
        past = self._lesson(yesterday, week="Н")
        future = self._lesson(tomorrow, week="Н")
        already = self._lesson(yesterday, start=time(19, 40), end=time(21, 10), week="+")
        resp = self.client.post(
            "/api/auto-attend/",
            data={"mark_all": True},
            content_type="application/json",
        )
        body = resp.json()
        self.assertTrue(body.get("ok"), body)
        self.assertEqual(body["marked"], 1)
        past.refresh_from_db()
        future.refresh_from_db()
        already.refresh_from_db()
        self.assertEqual(past.week_type, "+")
        self.assertEqual(future.week_type, "Н")
        self.assertEqual(already.week_type, "+")

        later = self._lesson(yesterday - timedelta(days=1), week="Н")
        tick_off = self.client.post(
            "/api/auto-attend/",
            data={"tick": True},
            content_type="application/json",
        )
        later.refresh_from_db()
        self.assertEqual(later.week_type, "Н")
        self.assertFalse(tick_off.json().get("auto_attend"))

        on = self.client.post(
            "/api/auto-attend/",
            data={"enabled": True},
            content_type="application/json",
        )
        self.assertTrue(on.json().get("auto_attend"))
        later.refresh_from_db()
        self.assertEqual(later.week_type, "+")

        extra = self._lesson(yesterday - timedelta(days=2), week="Н")
        tick_on = self.client.post(
            "/api/auto-attend/",
            data={"tick": True},
            content_type="application/json",
        )
        extra.refresh_from_db()
        self.assertEqual(extra.week_type, "+")
        self.assertGreaterEqual(tick_on.json()["marked"], 1)

    def test_copy_week_rejects_same_week(self):
        self._lesson(date(2026, 9, 7))
        resp = self.client.post(
            "/api/lessons/copy-week/",
            data={"from": "2026-09-07", "to": "2026-09-09"},
            content_type="application/json",
        )
        self.assertFalse(resp.json().get("ok"))
        self.assertEqual(resp.status_code, 400)


def _write_mirea_xlsx(path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "День недели"
    ws["B1"] = "№ пары"
    ws["C1"] = "Нач.занятий"
    ws["D1"] = "Оконч.занятий"
    ws["E1"] = "Неделя"
    ws["F1"] = "БСМО-31-25"
    ws["F2"] = "Предмет"
    ws["G2"] = "Вид занятий"
    ws["H2"] = "ФИО преподавателя"
    ws["I2"] = "№ ауд."
    ws["A3"] = "Понедельник"
    ws["B3"] = 1
    ws["C3"] = "18-00"
    ws["D3"] = "19-30"
    ws["E3"] = "I"
    ws["F3"] = "Облачные технологии"
    ws["G3"] = "пр"
    ws["H3"] = "Лебедев Александр Сергеевич"
    ws["I3"] = "А-402"
    ws["E4"] = "II"
    ws["A5"] = "Вторник"
    ws["B5"] = 2
    ws["C5"] = "19-40"
    ws["D5"] = "21-10"
    ws["E5"] = "I"
    ws["F5"] = "Иностранный язык кр. 1,3,5 н."
    ws["G5"] = "пр"
    ws["H5"] = "Орлова Елена Михайловна"
    ws["I5"] = "Б-201"
    ws["E6"] = "II"
    wb.save(path)


class ScheduleSyncTests(TestCase):
    def setUp(self):
        from journal.models import JournalSettings, StudentProfile

        self.user = User.objects.create_user(
            username="kononov.a.s@edu.mirea.ru",
            email="kononov.a.s@edu.mirea.ru",
            password="mirea2026",
        )
        self.group = Group.objects.create(slug="is", name="ИС-21-1", is_journal=True)
        self.bsmo = Group.objects.create(slug="bsmo", name="БСМО-31-25")
        StudentProfile.objects.create(
            user=self.user,
            short_name="Кононов А.",
            full_name="Кононов Артём Сергеевич",
            email="kononov.a.s@edu.mirea.ru",
            student_code="2561919",
            group=self.bsmo,
        )
        self.t1 = Teacher.objects.create(short_name="Лебедев А. С.", full_name="Лебедев Александр Сергеевич")
        self.subj = Subject.objects.create(name="Своя пара")
        JournalSettings.objects.create(
            pk=1,
            semester_start=date(2026, 9, 1),
            semester_end=date(2026, 9, 21),
        )
        self.client = Client()
        self.client.login(username="kononov.a.s@edu.mirea.ru", password="mirea2026")
        session = self.client.session
        session["schedule_edit"] = True
        session.save()

    def test_parse_and_sync_pending_accept_cancel(self):
        import tempfile
        from pathlib import Path

        from django.test import override_settings

        from journal.mirea_sync import parse_excel

        tmp = Path(tempfile.mkdtemp()) / "bsmo.xlsx"
        _write_mirea_xlsx(tmp)
        weekly = parse_excel(tmp, "БСМО-31-25")
        self.assertGreaterEqual(len(weekly), 2)
        names = {item.subject for item in weekly}
        self.assertIn("Облачные технологии", names)
        self.assertIn("Иностранный язык", names)

        with override_settings(MIREA_SCHEDULE_FILE=str(tmp), MIREA_SCHEDULE_GROUP="БСМО-31-25"):
            synced = self.client.post(
                "/api/schedule-sync/",
                data={"action": "sync"},
                content_type="application/json",
            )
        body = synced.json()
        self.assertTrue(body.get("ok"), body)
        self.assertGreater(body["pending"], 0)
        pending = Lesson.objects.filter(source="sync", accepted=False)
        self.assertEqual(pending.count(), body["pending"])
        self.assertTrue(pending.filter(date=date(2026, 9, 14), start_time=time(18, 0)).exists())
        self.assertTrue(pending.filter(date=date(2026, 9, 1), start_time=time(19, 40)).exists())

        ordinary = Client()
        ordinary.login(username="kononov.a.s@edu.mirea.ru", password="mirea2026")
        hidden = ordinary.get("/api/lessons/?journal=1").json()["lessons"]
        self.assertFalse(any(hidden.values()))

        shown = self.client.get("/api/lessons/?journal=1").json()["lessons"]
        self.assertTrue(shown.get("2026-09-14"))
        self.assertTrue(shown["2026-09-14"][0]["pending"])

        accepted = self.client.post(
            "/api/schedule-sync/",
            data={"action": "accept"},
            content_type="application/json",
        )
        self.assertTrue(accepted.json().get("ok"))
        self.assertEqual(Lesson.objects.filter(source="sync", accepted=False).count(), 0)
        visible = ordinary.get("/api/lessons/?journal=1").json()["lessons"]
        self.assertTrue(visible.get("2026-09-14"))
        self.assertEqual(visible["2026-09-14"][0]["source"], "sync")
        self.assertFalse(visible["2026-09-14"][0]["pending"])

        with override_settings(MIREA_SCHEDULE_FILE=str(tmp), MIREA_SCHEDULE_GROUP="БСМО-31-25"):
            self.client.post("/api/schedule-sync/", data={"action": "sync"}, content_type="application/json")
        self.assertEqual(Lesson.objects.filter(source="sync", accepted=False).count(), 0)

        cancel = self.client.post(
            "/api/schedule-sync/",
            data={"action": "cancel"},
            content_type="application/json",
        )
        self.assertEqual(cancel.json()["deleted"], 0)

    def test_sync_does_not_overwrite_manual(self):
        import tempfile
        from pathlib import Path

        from django.test import override_settings

        Lesson.objects.create(
            group=self.group,
            subject=self.subj,
            teacher=self.t1,
            date=date(2026, 9, 14),
            start_time=time(18, 0),
            end_time=time(19, 30),
            lesson_type="ПР",
            week_type="Н",
            room="999",
            source="manual",
            accepted=True,
        )
        tmp = Path(tempfile.mkdtemp()) / "bsmo.xlsx"
        _write_mirea_xlsx(tmp)
        with override_settings(MIREA_SCHEDULE_FILE=str(tmp), MIREA_SCHEDULE_GROUP="БСМО-31-25"):
            body = self.client.post(
                "/api/schedule-sync/",
                data={"action": "sync"},
                content_type="application/json",
            ).json()
        self.assertTrue(body.get("ok"), body)
        self.assertGreaterEqual(body["skipped_manual"], 1)
        kept = Lesson.objects.get(date=date(2026, 9, 14), start_time=time(18, 0))
        self.assertEqual(kept.source, "manual")
        self.assertEqual(kept.room, "999")
        self.assertEqual(kept.subject_id, self.subj.id)

        canceled = self.client.post(
            "/api/schedule-sync/",
            data={"action": "cancel"},
            content_type="application/json",
        ).json()
        self.assertGreater(canceled["deleted"], 0)
        self.assertTrue(Lesson.objects.filter(pk=kept.pk).exists())
